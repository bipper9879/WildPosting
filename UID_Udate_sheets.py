import argparse
import glob
import os
import re
import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl
import pandas as pd

import uid_utils

BASE_LOCAL_PATH = uid_utils.BASE_LOCAL_PATH
REGISTRY_PATH = uid_utils.REGISTRY_PATH

CITIES_CONFIG = {
    "philly": {"folder": "philly", "pattern": "philly"},
    "dc": {"folder": "dc", "pattern": "dc"},
    "boston": {"folder": "boston", "pattern": "boston"},
}


def _normalize_city_name(city_name):
    if not city_name:
        return None
    key = city_name.strip().lower()
    if key in {"phl", "philly", "philadelphia"}:
        return "philly"
    if key in {"dc", "d.c.", "washington"}:
        return "dc"
    if key in {"bos", "boston"}:
        return "boston"
    return None


def _matches_skip_pattern(file_name, skip_patterns):
    lower_name = file_name.lower()
    return any(pattern.lower() in lower_name for pattern in skip_patterns)


def _is_dated_workbook(file_name, city_pattern):
    if not file_name.lower().endswith(".xlsx"):
        return False
    if file_name.startswith("~$"):
        return False
    if file_name.lower().endswith("_workbook.xlsx"):
        return False
    if "country_uid_registry" in file_name.lower():
        return False
    if city_pattern not in file_name.lower():
        return False

    date_pattern = re.compile(r"(\d{2}-\d{2}-\d{4}|\d{4}-\d{2}-\d{2})")
    return bool(date_pattern.search(file_name))


def discover_dated_workbooks(city_directory, city_pattern, skip_patterns):
    candidates = []
    for filepath in glob.glob(os.path.join(city_directory, "*.xlsx")):
        file_name = os.path.basename(filepath)
        if not _is_dated_workbook(file_name, city_pattern):
            continue
        if _matches_skip_pattern(file_name, skip_patterns):
            continue
        candidates.append(filepath)
    return sorted(candidates, key=os.path.getmtime, reverse=True)


def update_single_workbook(file_path, uid_lookup):
    wb = openpyxl.load_workbook(file_path)
    file_updated = False

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            location_col_idx = None
            street_view_col_idx = None

            for col in range(1, ws.max_column + 1):
                header_val = str(ws.cell(row=1, column=col).value).strip()
                normalized_header = header_val.lower()

                if header_val == "Location":
                    location_col_idx = col
                if normalized_header in {"google street view", "street view", "streetview"}:
                    street_view_col_idx = col

            if street_view_col_idx is None:
                street_view_col_idx = 6

            if not location_col_idx:
                continue

            if ws.cell(row=1, column=1).value != "Site Code":
                ws.cell(row=1, column=1).value = "Site Code"
                file_updated = True
            if ws.cell(row=1, column=7).value != "Lat":
                ws.cell(row=1, column=7).value = "Lat"
                file_updated = True
            if ws.cell(row=1, column=8).value != "Lon":
                ws.cell(row=1, column=8).value = "Lon"
                file_updated = True

            for row_idx in range(2, ws.max_row + 1):
                raw_loc = ws.cell(row=row_idx, column=location_col_idx).value
                if raw_loc is not None:
                    location_name = str(raw_loc).strip()
                    if location_name in uid_lookup:
                        target_uid = uid_lookup[location_name]
                        cell_a = ws.cell(row=row_idx, column=1)

                        if str(cell_a.value).strip() != target_uid:
                            cell_a.value = target_uid
                            file_updated = True

                cell_f = ws.cell(row=row_idx, column=street_view_col_idx)
                cell_f_url = cell_f.hyperlink.target if cell_f.hyperlink else cell_f.value

                if cell_f_url:
                    lat, lon = uid_utils.extract_gps_from_url(str(cell_f_url))
                    if lat and lon:
                        cell_g_lat = ws.cell(row=row_idx, column=7)
                        cell_h_lon = ws.cell(row=row_idx, column=8)

                        if str(cell_g_lat.value) != str(lat) or str(cell_h_lon.value) != str(lon):
                            cell_g_lat.value = float(lat)
                            cell_h_lon.value = float(lon)
                            file_updated = True

        if file_updated:
            wb.save(file_path)
    finally:
        wb.close()

    return file_updated


def process_selected_cities(selected_cities, all_dated=False, skip_patterns=None):
    skip_patterns = skip_patterns or []

    if not os.path.exists(REGISTRY_PATH):
        return [f"❌ Central registry file not found: {REGISTRY_PATH}"]

    reg_df = pd.read_excel(REGISTRY_PATH, sheet_name="Master")
    reg_df["Location Name"] = reg_df["Location Name"].astype(str).str.strip()
    uid_lookup = dict(zip(reg_df["Location Name"], reg_df["UID"]))

    summary_log = []
    for city in selected_cities:
        config = CITIES_CONFIG[city]
        city_dir = os.path.join(BASE_LOCAL_PATH, config["folder"])

        if not os.path.exists(city_dir):
            summary_log.append(f"❌ {city}: Folder missing ({city_dir})")
            continue

        candidates = discover_dated_workbooks(city_dir, config["pattern"], skip_patterns)
        if not candidates:
            summary_log.append(f"⚠️ {city}: No matching dated files found in {city_dir}")
            continue

        workbooks_to_process = candidates if all_dated else candidates[:1]
        updated_count = 0

        for target_workbook in workbooks_to_process:
            file_name = os.path.basename(target_workbook)
            if not uid_utils.ensure_file_is_free(target_workbook, file_name):
                summary_log.append(f"❌ {city}: File locked, skipped -> {file_name}")
                continue

            try:
                was_updated = update_single_workbook(target_workbook, uid_lookup)
                if was_updated:
                    updated_count += 1
                    summary_log.append(f"✅ {city}: Updated -> {file_name}")
                else:
                    summary_log.append(f"⏭️ {city}: No changes -> {file_name}")
            except Exception as exc:
                summary_log.append(f"❌ {city}: Error in {file_name}: {exc}")

        summary_log.append(
            f"ℹ️ {city}: Processed {len(workbooks_to_process)} workbook(s), updated {updated_count}."
        )

    return summary_log


def launch_gui():
    root = tk.Tk()
    root.title("City Sync Selector")
    root.geometry("360x280")
    root.resizable(False, False)
    root.eval("tk::PlaceWindow . center")

    label = ttk.Label(root, text="Select cities to sync values:", font=("Arial", 11, "bold"))
    label.pack(pady=12)

    checkbox_vars = {}
    frame = ttk.Frame(root)
    frame.pack(pady=5)

    for city in CITIES_CONFIG.keys():
        var = tk.BooleanVar(value=True)
        checkbox_vars[city] = var
        chk = ttk.Checkbutton(frame, text=city.title(), variable=var)
        chk.pack(anchor="w", pady=4)

    all_dated_var = tk.BooleanVar(value=True)
    ttk.Checkbutton(
        root,
        text="Process all dated workbooks (not just latest)",
        variable=all_dated_var,
    ).pack(pady=(8, 0))

    def on_submit():
        chosen = [city for city, var in checkbox_vars.items() if var.get()]
        if not chosen:
            messagebox.showwarning("Selection Empty", "Please check at least one city box.")
            return
        root.destroy()
        lines = process_selected_cities(chosen, all_dated=all_dated_var.get(), skip_patterns=[])
        messagebox.showinfo("Sync Complete", "\n".join(lines))

    submit_btn = ttk.Button(root, text="Run Sync Engine", command=on_submit)
    submit_btn.pack(pady=15)
    root.mainloop()


def parse_args():
    parser = argparse.ArgumentParser(description="Sync UID + Lat/Lon values into dated city workbooks.")
    parser.add_argument(
        "--cities",
        nargs="+",
        default=["philly", "dc", "boston"],
        help="City names/codes to process (e.g. philly dc boston)",
    )
    parser.add_argument(
        "--all-dated",
        action="store_true",
        help="Process all dated workbooks for each city (default CLI behavior if no --latest-only).",
    )
    parser.add_argument(
        "--latest-only",
        action="store_true",
        help="Process only the latest dated workbook per city.",
    )
    parser.add_argument(
        "--skip-patterns",
        nargs="*",
        default=[],
        help="Filename substrings to skip (e.g. 7-22 07-22).",
    )
    parser.add_argument(
        "--gui",
        action="store_true",
        help="Launch the interactive city selection GUI.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()

    if args.gui:
        launch_gui()
    else:
        normalized_cities = []
        for city_arg in args.cities:
            normalized = _normalize_city_name(city_arg)
            if normalized and normalized not in normalized_cities:
                normalized_cities.append(normalized)

        if not normalized_cities:
            normalized_cities = ["philly", "dc", "boston"]

        process_all = True
        if args.latest_only:
            process_all = False
        elif args.all_dated:
            process_all = True

        output_lines = process_selected_cities(
            normalized_cities,
            all_dated=process_all,
            skip_patterns=args.skip_patterns,
        )
        for line in output_lines:
            print(line)
