import os
import openpyxl
import re

# Graphical User Interface Libraries
import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry

# ==================== CONFIGURATION ====================
BASE_POSTERS_DIR = r"C:\Users\bippe\OneDrive\Documents\posters"
# =======================================================

def extract_coords_from_streetview(url):
    """Extracts latitude and longitude from a Google Street View URL."""
    if not url:
        return None
    match = re.search(r'([-+]?\d+\.\d+),\s*([-+]?\d+\.\d+)', url)
    if match:
        return f"{match.group(1)},{match.group(2)}"
    return None

def process_submission():
    city = city_entry.get().strip()
    selected_date = cal.get_date()
    
    if not city:
        messagebox.showwarning("Input Error", "Please enter a city name.")
        return
        
    date_str = selected_date.strftime("%m-%d-%Y")
    workbook_name = f"{city}_Workbook.xlsx"
    excel_path = os.path.join(BASE_POSTERS_DIR, city, workbook_name)
    
    if not os.path.exists(excel_path):
        status_label.config(text=f"Error: {workbook_name} not found!", foreground="red")
        return

    status_label.config(text="Processing safe links...", foreground="blue")
    root.update_idletasks()

    try:
        wb = openpyxl.load_workbook(excel_path)
        
        if date_str not in wb.sheetnames:
            status_label.config(text=f"Error: Sheet '{date_str}' not found!", foreground="red")
            wb.close()
            return
            
        sheet = wb[date_str]
        
        # --- RESET AND PURGE CELLS ---
        # Wipes out text values, formulas, and old styles completely
        for cell_id in ["G2", "G3"]:
            sheet[cell_id].value = None
            sheet[cell_id].hyperlink = None
            sheet[cell_id].style = 'Normal'
            sheet[cell_id].number_format = 'General'
        
        coordinates_list = []
        
        # Read from row 2 downward to extract Column F's coordinates
        for row in range(2, sheet.max_row + 1):
            sv_cell = sheet.cell(row=row, column=6) # Column F
            sv_url = sv_cell.hyperlink.target if sv_cell.hyperlink else sv_cell.value
            
            if sv_url:
                coords = extract_coords_from_streetview(str(sv_url))
                if coords and coords not in coordinates_list:
                    coordinates_list.append(coords)

        total_points = len(coordinates_list)

        if total_points == 0:
            status_label.config(text="No valid coordinates found in Column F.", foreground="red")
            wb.close()
            return

        base_dir_url = "https://www.google.com/maps/dir/"
        
        # --- BYPASSING 255-CHAR LIMIT VIA METADATA OBJECTS ---
        # Instead of =HYPERLINK(), we drop the raw URL string into the sheet's XML structure.
        # This allows links up to 2084 characters to function seamlessly.
        
        if total_points <= 20:
            path_string = "/".join(coordinates_list)
            final_url = f"{base_dir_url}{path_string}"
            
            # Write text label, map underlying URL object, and color text blue
            sheet["G2"].value = f"Pins (1-{total_points})"
            sheet["G2"].hyperlink = final_url
            sheet["G2"].style = "Hyperlink"
            status_msg = f"Success! All {total_points} items mapped to G2."
        else:
            midpoint = total_points // 2
            cluster_1 = coordinates_list[:midpoint]
            cluster_2 = coordinates_list[midpoint:]
            
            url_1 = f"{base_dir_url}{'/'.join(cluster_1)}"
            url_2 = f"{base_dir_url}{'/'.join(cluster_2)}"
            
            # Bind Cluster 1 to G2
            sheet["G2"].value = f"Pins (1-{midpoint})"
            sheet["G2"].hyperlink = url_1
            sheet["G2"].style = "Hyperlink"
            
            # Bind Cluster 2 to G3
            sheet["G3"].value = f"Pins ({midpoint + 1}-{total_points})"
            sheet["G3"].hyperlink = url_2
            sheet["G3"].style = "Hyperlink"
            status_msg = f"Split {total_points} items between G2 and G3!"

        wb.save(excel_path)
        status_label.config(text=status_msg, foreground="green")
        city_entry.delete(0, tk.END)

    except PermissionError:
        status_label.config(text="Error: Close the Excel file first!", foreground="red")
    except Exception as e:
        status_label.config(text=f"Error: {str(e)}", foreground="red")

# Create UI Window Layout
root = tk.Tk()
root.title("Column G Multi-Pin Linker")
root.geometry("400x280")

ttk.Label(root, text="Enter City Name (e.g., Philly):", font=("Arial", 10)).pack(pady=(15, 2))
city_entry = ttk.Entry(root, width=32)
city_entry.pack(pady=5)
city_entry.focus()

ttk.Label(root, text="Select Sheet Date:", font=("Arial", 10)).pack(pady=(10, 2))
cal = DateEntry(root, width=29, background='darkblue', foreground='white', borderwidth=2, date_pattern='mm-dd-yyyy')
cal.pack(pady=5)

submit_btn = ttk.Button(root, text="Generate Stacked Map Links", command=process_submission)
submit_btn.pack(pady=15)

status_label = ttk.Label(root, text="Ready", font=("Arial", 9, "italic"))
status_label.pack(pady=5)

root.mainloop()
