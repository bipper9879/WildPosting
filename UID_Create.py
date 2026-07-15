import os
import sys
import ctypes
import pandas as pd
import openpyxl

# Import your custom utility modules from Part 1
import uid_utils
from uid_utils import log_message, show_popup, ensure_file_is_free, create_text_fingerprint, extract_gps_from_url

# --- SYSTEM-WIDE EMERGENCY CRASH TRAP ---
def crash_notifier(exctype, value, tb):
    import traceback
    error_msg = "".join(traceback.format_exception(exctype, value, tb))
    try:
        with open(uid_utils.LOG_PATH, "a", encoding="utf-8") as lf:
            lf.write(f"[FATAL SCRIPT CRASH]\n{error_msg}\n")
    except Exception:
        pass
    ctypes.windll.user32.MessageBoxW(0, f"Critical Sync Engine Crash:\n\n{error_msg}\n\nReview UID_Automation_Log.txt!", "Script Crashed!", 0x10 | 0x00)
    sys.__excepthook__(exctype, value, tb)

sys.excepthook = crash_notifier


def run_complete_sync():
    uid_utils.init_system_log()
    log_message("🚀 Starting Multi-City Synchronization Engine...")
    
    target_cities = ["philly", "dc", "boston"]
    registry_data = []
    
    # --- PHASE 1: LOAD AND INDEX CITY WORKBOOKS (READ-ONLY) ---
    for city_folder in target_cities:
        city_dir_path = os.path.join(uid_utils.BASE_LOCAL_PATH, city_folder)
        if not os.path.exists(city_dir_path): continue
            
        workbook_name = f"{city_folder}_Workbook.xlsx"
        file_path = os.path.join(city_dir_path, workbook_name)
        if not os.path.exists(file_path): continue
            
        city_prefix = city_folder[:3].upper()
        if not ensure_file_is_free(file_path, workbook_name): return
            
        try:
            df = pd.read_excel(file_path, sheet_name='Master', header=0)
            df.columns = df.columns.str.strip()
            log_message(f"Successfully loaded {workbook_name} [Sheet: Master]")
            
            if 'Location' not in df.columns:
                log_message(f"❌ Error: Could not find 'Location' column header inside {workbook_name}.")
                continue
                
            for index, row in df.iterrows():
                location_name = str(row['Location']).strip()
                if location_name == "" or pd.isna(row['Location']) or location_name.lower() in ["nan", "location"]:
                    continue
                    
                spot_hash = create_text_fingerprint(location_name)
                final_uid = f"{city_prefix}-{spot_hash}"
                registry_data.append({
                    "UID": final_uid,
                    "City": city_prefix,
                    "Location Name": location_name,
                    "Source Workbook": workbook_name
                })
        except Exception as e:
            log_message(f"❌ Critical failure reading {workbook_name}: {str(e)}")
            return

    if not registry_data:
        log_message("❌ Execution Stopped: No data rows found inside your 'Location' columns.")
        show_popup("Sync Stopped", "No text records were found inside your 'Location' columns yet.")
        return

    # --- PHASE 2: WRITE CENTRAL MASTER SHEET ---
    if not ensure_file_is_free(uid_utils.REGISTRY_PATH, "Country_UID_Registry.xlsx"): return
        
    try:
        master_df = pd.DataFrame(registry_data)
        master_df.drop_duplicates(subset=["UID"], keep="first", inplace=True)
        with pd.ExcelWriter(uid_utils.REGISTRY_PATH, engine='openpyxl') as writer:
            master_df.to_excel(writer, sheet_name='Master', index=False)
        log_message("Vault Saved: Country_UID_Registry.xlsx successfully populated.")
    except Exception as e:
        log_message(f"❌ Failed saving registry file matrix: {str(e)}")
        return

    # --- PHASE 3: SURGICAL FORMAT-SAFE CELL INJECTION (UID + HEADERS + AUTO-GPS) ---
    log_message("Starting Phase 3: Stamping UIDs, forcing headers, and auto-extracting GPS coordinates...")
    uid_lookup = dict(zip(master_df["Location Name"], master_df["UID"]))
    
    # Track any locations that are missing Street View hyperlinks
    missing_links_report = []
    
    for city_folder in target_cities:
        city_dir_path = os.path.join(uid_utils.BASE_LOCAL_PATH, city_folder)
        workbook_name = f"{city_folder}_Workbook.xlsx"
        file_path = os.path.join(city_dir_path, workbook_name)
        if not os.path.exists(file_path): continue
        if not ensure_file_is_free(file_path, workbook_name): return
            
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb['Master']
            
            location_col_idx = None
            street_view_col_idx = None
            
            for col in range(1, ws.max_column + 1):
                header_val = str(ws.cell(row=1, column=col).value).strip()
                normalized_header = header_val.lower()
                if header_val == 'Location':
                    location_col_idx = col
                if normalized_header in {'google street view', 'street view', 'streetview'}:
                    street_view_col_idx = col

            if street_view_col_idx is None:
                street_view_col_idx = 6
            
            if not location_col_idx:
                log_message(f"❌ Could not find 'Location' column header in {workbook_name}. Skipping file.")
                wb.close()
                continue
            
            ws.cell(row=1, column=1).value = "Site Code"
            ws.cell(row=1, column=7).value = "Lat"
            ws.cell(row=1, column=8).value = "Lon"
            
            changes_made = True
            
            for r in range(2, ws.max_row + 1):
                raw_loc = ws.cell(row=r, column=location_col_idx).value
                if raw_loc is None: continue
                
                location_name = str(raw_loc).strip()
                if location_name in uid_lookup:
                    target_uid = uid_lookup[location_name]
                    
                    # Stamp UID safely as text
                    cell_a = ws.cell(row=r, column=1)
                    cell_a.number_format = '@' 
                    if str(cell_a.value).strip() != target_uid:
                        cell_a.value = target_uid
                        changes_made = True
                    
                    # Check cell and hyperlinks for an active Street View URL
                    cell_f_url = ws.cell(row=r, column=street_view_col_idx).value
                    if ws.cell(row=r, column=street_view_col_idx).hyperlink:
                        cell_f_url = ws.cell(row=r, column=street_view_col_idx).hyperlink.target
                        
                    lat, lon = extract_gps_from_url(cell_f_url)
                    
                    if lat and lon:
                        cell_g_lat = ws.cell(row=r, column=7)
                        cell_h_lon = ws.cell(row=r, column=8)
                        
                        if str(cell_g_lat.value) != str(lat) or str(cell_h_lon.value) != str(lon):
                            cell_g_lat.value = float(lat)
                            cell_h_lon.value = float(lon)
                            changes_made = True
                    else:
                        # Log the location if Street View data is missing or empty
                        missing_links_report.append(f"[{city_folder.upper()}] Row {r}: {location_name}")
            
            if changes_made:
                wb.save(file_path)
                log_message(f"Success: Preserved formatting, updated headers, and extracted GPS into {workbook_name}.")
            else:
                log_message(f"Checked {workbook_name}: Already up to date.")
            wb.close()
        except Exception as e:
            log_message(f"❌ Failed processing on {workbook_name}: {str(e)}")
            return

    log_message("🎉 All processing tasks finished cleanly. Multi-city network synchronized.")
    
    # --- PHASE 4: FINAL NOTIFICATION LOGIC WITH MISSING LINKS REPORT ---
    if missing_links_report:
        alert_msg = "All systems synced successfully!\n\n⚠️ NOTE: The following locations are missing valid Street View links (Lat/Lon skipped):\n\n"
        alert_msg += "\n".join(missing_links_report[:12]) # Show the first 12 lines in pop-up window
        
        if len(missing_links_report) > 12:
            alert_msg += f"\n...and {len(missing_links_report) - 12} more locations. See text log file!"
            
        # Log the complete list for easy review
        log_message("⚠️ MISSING STREET VIEW LINKS LEDGER:")
        for missing_item in missing_links_report:
            log_message(f"   Missing link details: {missing_item}")
            
        show_popup("Sync Complete (With Alerts)", alert_msg, is_error=False)
    else:
        show_popup("Sync Complete!", "All systems synced successfully!\n\nEvery location matched a valid Street View URL and generated GPS values.", is_error=False)

if __name__ == "__main__":
    run_complete_sync()
